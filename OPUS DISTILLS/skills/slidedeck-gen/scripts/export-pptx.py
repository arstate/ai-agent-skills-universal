#!/usr/bin/env python3
"""
export-pptx.py — Autonomous 16:9 PowerPoint (.pptx) Exporter for slidedeck-gen
Generates genuine, editable PowerPoint presentations with native shapes, text boxes,
metric cards, and theme palettes using python-pptx.
Zero npm dependencies required.
"""

import sys
import os
import re
import json
from bs4 import BeautifulSoup
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

THEME_PALETTES = {
    "bold-signal": {
        "bg": RGBColor(24, 24, 27),
        "card_bg": RGBColor(36, 36, 40),
        "card_border": RGBColor(63, 63, 70),
        "title": RGBColor(255, 255, 255),
        "body": RGBColor(244, 244, 245),
        "muted": RGBColor(161, 161, 170),
        "accent": RGBColor(255, 87, 34),  # High-impact electric orange
    },
    "electric-studio": {
        "bg": RGBColor(14, 17, 24),
        "card_bg": RGBColor(20, 24, 36),
        "card_border": RGBColor(30, 38, 56),
        "title": RGBColor(255, 255, 255),
        "body": RGBColor(226, 232, 240),
        "muted": RGBColor(148, 163, 184),
        "accent": RGBColor(59, 130, 246),  # Cobalt electric blue
    },
    "creative-voltage": {
        "bg": RGBColor(14, 9, 43),
        "card_bg": RGBColor(26, 18, 64),
        "card_border": RGBColor(46, 33, 105),
        "title": RGBColor(255, 255, 255),
        "body": RGBColor(241, 240, 251),
        "muted": RGBColor(139, 133, 184),
        "accent": RGBColor(168, 85, 247),  # Neon violet
    },
    "editorial-forest": {
        "bg": RGBColor(14, 36, 27),
        "card_bg": RGBColor(19, 48, 36),
        "card_border": RGBColor(29, 72, 55),
        "title": RGBColor(255, 255, 255),
        "body": RGBColor(244, 237, 226),
        "muted": RGBColor(197, 214, 204),
        "accent": RGBColor(212, 175, 55),  # Champagne gold
    },
    "emerald-editorial": {
        "bg": RGBColor(10, 54, 41),
        "card_bg": RGBColor(14, 68, 52),
        "card_border": RGBColor(21, 90, 69),
        "title": RGBColor(255, 255, 255),
        "body": RGBColor(232, 245, 241),
        "muted": RGBColor(184, 224, 212),
        "accent": RGBColor(16, 185, 129),  # Mint emerald
    },
    "neo-grid-bold": {
        "bg": RGBColor(255, 255, 255),
        "card_bg": RGBColor(250, 250, 250),
        "card_border": RGBColor(24, 24, 27),
        "title": RGBColor(9, 9, 11),
        "body": RGBColor(24, 24, 27),
        "muted": RGBColor(113, 113, 122),
        "accent": RGBColor(0, 0, 0),  # Pure stark black
    },
    "cobalt-cad": {
        "bg": RGBColor(0, 24, 69),
        "card_bg": RGBColor(2, 35, 91),
        "card_border": RGBColor(4, 102, 200),
        "title": RGBColor(255, 255, 255),
        "body": RGBColor(192, 199, 214),
        "muted": RGBColor(125, 133, 151),
        "accent": RGBColor(0, 212, 255),  # Blueprint cyan
    },
    "capsule-modern": {
        "bg": RGBColor(30, 41, 59),
        "card_bg": RGBColor(39, 53, 79),
        "card_border": RGBColor(51, 65, 85),
        "title": RGBColor(255, 255, 255),
        "body": RGBColor(248, 250, 252),
        "muted": RGBColor(148, 163, 184),
        "accent": RGBColor(56, 189, 248),  # Sky blue
    },
    "soft-editorial": {
        "bg": RGBColor(250, 247, 242),
        "card_bg": RGBColor(255, 255, 255),
        "card_border": RGBColor(230, 223, 211),
        "title": RGBColor(26, 24, 22),
        "body": RGBColor(42, 39, 35),
        "muted": RGBColor(122, 115, 104),
        "accent": RGBColor(139, 94, 60),  # Warm bronze
    },
    "monochrome-minimal": {
        "bg": RGBColor(10, 10, 10),
        "card_bg": RGBColor(18, 18, 18),
        "card_border": RGBColor(36, 36, 36),
        "title": RGBColor(255, 255, 255),
        "body": RGBColor(255, 255, 255),
        "muted": RGBColor(136, 136, 136),
        "accent": RGBColor(255, 255, 255),
    },
    "broadside-poster": {
        "bg": RGBColor(39, 39, 42),
        "card_bg": RGBColor(63, 63, 70),
        "card_border": RGBColor(82, 82, 91),
        "title": RGBColor(255, 255, 255),
        "body": RGBColor(244, 244, 245),
        "muted": RGBColor(161, 161, 170),
        "accent": RGBColor(225, 29, 72),  # Bold rose red
    },
    "cyber-glow": {
        "bg": RGBColor(11, 15, 25),
        "card_bg": RGBColor(20, 28, 44),
        "card_border": RGBColor(51, 65, 85),
        "title": RGBColor(248, 250, 252),
        "body": RGBColor(203, 213, 225),
        "muted": RGBColor(148, 163, 184),
        "accent": RGBColor(16, 185, 129),
    },
    "tech-code": {
        "bg": RGBColor(24, 25, 30),
        "card_bg": RGBColor(32, 34, 42),
        "card_border": RGBColor(46, 50, 61),
        "title": RGBColor(255, 255, 255),
        "body": RGBColor(209, 213, 219),
        "muted": RGBColor(156, 163, 175),
        "accent": RGBColor(56, 189, 248),
    },
    "noir-luxe": {
        "bg": RGBColor(7, 7, 7),
        "card_bg": RGBColor(17, 17, 17),
        "card_border": RGBColor(38, 38, 38),
        "title": RGBColor(255, 255, 255),
        "body": RGBColor(229, 229, 229),
        "muted": RGBColor(163, 163, 163),
        "accent": RGBColor(212, 175, 55),
    },
    "high-growth": {
        "bg": RGBColor(255, 255, 255),
        "card_bg": RGBColor(250, 250, 250),
        "card_border": RGBColor(228, 228, 231),
        "title": RGBColor(9, 9, 11),
        "body": RGBColor(63, 63, 70),
        "muted": RGBColor(113, 113, 122),
        "accent": RGBColor(29, 78, 216),
    },
    "modern-clean": {
        "bg": RGBColor(255, 255, 255),
        "card_bg": RGBColor(250, 250, 250),
        "card_border": RGBColor(228, 228, 231),
        "title": RGBColor(24, 24, 27),
        "body": RGBColor(63, 63, 70),
        "muted": RGBColor(113, 113, 122),
        "accent": RGBColor(2, 132, 199),
    }
}

# Alias theme numbers
for i in range(1, 13):
    key = f"theme{i:02d}"
    if i in (2, 6, 12):
        THEME_PALETTES[key] = THEME_PALETTES["cyber-glow"]
    elif i in (3, 10):
        THEME_PALETTES[key] = THEME_PALETTES["tech-code"]
    elif i == 8:
        THEME_PALETTES[key] = THEME_PALETTES["noir-luxe"]
    elif i in (5, 11):
        THEME_PALETTES[key] = THEME_PALETTES["high-growth"]
    else:
        THEME_PALETTES[key] = THEME_PALETTES["modern-clean"]


def parse_slides_from_html(html_content):
    soup = BeautifulSoup(html_content, "html.parser")
    slides_data = []

    html_el = soup.find("html")
    theme_attr = html_el.get("data-theme", "modern-clean") if html_el else "modern-clean"

    slide_els = soup.find_all("section", class_="slide")
    for idx, slide_el in enumerate(slide_els):
        classes = slide_el.get("class", [])
        layout = "standard"
        for c in classes:
            if c.startswith("layout-"):
                layout = c.replace("layout-", "")
                break

        kicker_el = slide_el.find(class_="kicker")
        kicker = kicker_el.get_text(strip=True) if kicker_el else ""

        h1_el = slide_el.find("h1")
        title = h1_el.get_text(strip=True) if h1_el else ""

        h2_el = slide_el.find("h2")
        h3_el = slide_el.find("h3")
        subtitle = h2_el.get_text(strip=True) if h2_el else (h3_el.get_text(strip=True) if h3_el else "")

        # Extract quote
        quote_el = slide_el.find("blockquote")
        quote = quote_el.get_text(strip=True) if quote_el else ""
        quote_attr_el = slide_el.find(class_="quote-attribution")
        quote_attr = quote_attr_el.get_text(strip=True) if quote_attr_el else ""

        # Extract metric cards
        metrics = []
        for m_card in slide_el.find_all(class_="metric-card"):
            val_el = m_card.find(class_="metric-value")
            lbl_el = m_card.find(class_="metric-label")
            desc_el = m_card.find(class_="metric-desc")
            metrics.append({
                "value": val_el.get_text(strip=True) if val_el else "",
                "label": lbl_el.get_text(strip=True) if lbl_el else "",
                "desc": desc_el.get_text(strip=True) if desc_el else ""
            })

        # Extract grid cards
        cards = []
        for g_card in slide_el.find_all(class_="grid-card"):
            t_el = g_card.find(["h4", "h3", "h5"])
            p_el = g_card.find("p")
            cards.append({
                "title": t_el.get_text(strip=True) if t_el else "",
                "desc": p_el.get_text(strip=True) if p_el else ""
            })

        # Extract flow steps
        steps = []
        for f_step in slide_el.find_all(class_="flow-step"):
            num_el = f_step.find(class_="step-num")
            t_el = f_step.find(["h4", "h3"])
            p_el = f_step.find("p")
            steps.append({
                "num": num_el.get_text(strip=True) if num_el else "",
                "title": t_el.get_text(strip=True) if t_el else "",
                "desc": p_el.get_text(strip=True) if p_el else ""
            })

        # Extract bullet points
        bullets = []
        for li in slide_el.find_all("li"):
            bullets.append(li.get_text(strip=True))

        # Extract paragraphs (excluding cards)
        paragraphs = []
        for p in slide_el.find_all("p", recursive=False):
            t = p.get_text(strip=True)
            if t and t != subtitle:
                paragraphs.append(t)

        slides_data.append({
            "index": idx + 1,
            "layout": layout,
            "kicker": kicker,
            "title": title,
            "subtitle": subtitle,
            "quote": quote,
            "quote_attr": quote_attr,
            "metrics": metrics,
            "cards": cards,
            "steps": steps,
            "bullets": bullets,
            "paragraphs": paragraphs
        })

    return theme_attr, slides_data


def export_deck_to_pptx(html_path, out_pptx_path):
    with open(html_path, "r", encoding="utf-8") as f:
        html_content = f.read()

    theme_name, slides_data = parse_slides_from_html(html_content)
    if not slides_data:
        raise ValueError(f"No valid <section class=\"slide\"> elements found in {html_path}. Ensure deck is properly compiled.")
    palette = THEME_PALETTES.get(theme_name.lower(), THEME_PALETTES["modern-clean"])

    prs = pptx.Presentation()
    # 16:9 Widescreen Presentation (13.333 x 7.5 inches)
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    for s in slides_data:
        slide = prs.slides.add_slide(blank_layout)

        # Slide background fill
        bg_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg_shape.fill.solid()
        bg_shape.fill.fore_color.rgb = palette["bg"]
        bg_shape.line.fill.background()

        # Cover Slide Layout
        if s["layout"] == "cover":
            top = Inches(2.0)
            if s["kicker"]:
                kbox = slide.shapes.add_textbox(Inches(1.2), top, Inches(10.9), Inches(0.4))
                ktf = kbox.text_frame
                ktf.word_wrap = True
                kp = ktf.paragraphs[0]
                kp.text = s["kicker"].upper()
                kp.font.size = Pt(13)
                kp.font.bold = True
                kp.font.color.rgb = palette["accent"]
                top += Inches(0.5)

            tbox = slide.shapes.add_textbox(Inches(1.2), top, Inches(10.9), Inches(2.0))
            ttf = tbox.text_frame
            ttf.word_wrap = True
            tp = ttf.paragraphs[0]
            tp.text = s["title"] or "Presentation Title"
            tp.font.size = Pt(44)
            tp.font.bold = True
            tp.font.color.rgb = palette["title"]

            top += Inches(1.8)
            if s["subtitle"]:
                sbox = slide.shapes.add_textbox(Inches(1.2), top, Inches(10.9), Inches(1.5))
                stf = sbox.text_frame
                stf.word_wrap = True
                sp = stf.paragraphs[0]
                sp.text = s["subtitle"]
                sp.font.size = Pt(20)
                sp.font.color.rgb = palette["muted"]

        # Quote Spotlight Layout
        elif s["layout"] == "quote" or (s["quote"] and not s["metrics"] and not s["cards"]):
            curr_top = Inches(2.0)
            if s["kicker"]:
                kbox = slide.shapes.add_textbox(Inches(1.5), curr_top, Inches(10.333), Inches(0.4))
                ktf = kbox.text_frame
                ktf.word_wrap = True
                kp = ktf.paragraphs[0]
                kp.alignment = PP_ALIGN.CENTER
                kp.text = s["kicker"].upper()
                kp.font.size = Pt(12)
                kp.font.bold = True
                kp.font.color.rgb = palette["accent"]
                curr_top += Inches(0.6)

            qbox = slide.shapes.add_textbox(Inches(1.5), curr_top, Inches(10.333), Inches(2.8))
            qtf = qbox.text_frame
            qtf.word_wrap = True
            qp = qtf.paragraphs[0]
            qp.alignment = PP_ALIGN.CENTER
            qp.text = f"“{s['quote'] or s['title']}”"
            qp.font.size = Pt(30)
            qp.font.italic = True
            qp.font.color.rgb = palette["title"]

            curr_top += Inches(3.0)
            if s["quote_attr"] or s["subtitle"]:
                abox = slide.shapes.add_textbox(Inches(1.5), curr_top, Inches(10.333), Inches(0.8))
                atf = abox.text_frame
                atf.word_wrap = True
                ap = atf.paragraphs[0]
                ap.alignment = PP_ALIGN.CENTER
                ap.text = f"— {s['quote_attr'] or s['subtitle']}"
                ap.font.size = Pt(16)
                ap.font.bold = True
                ap.font.color.rgb = palette["accent"]

        else:
            # Standard & Content Slides
            curr_top = Inches(0.8)

            if s["kicker"]:
                kbox = slide.shapes.add_textbox(Inches(1.0), curr_top, Inches(11.333), Inches(0.35))
                ktf = kbox.text_frame
                ktf.word_wrap = True
                kp = ktf.paragraphs[0]
                kp.text = s["kicker"].upper()
                kp.font.size = Pt(11)
                kp.font.bold = True
                kp.font.color.rgb = palette["accent"]
                curr_top += Inches(0.4)

            # Slide Title
            tbox = slide.shapes.add_textbox(Inches(1.0), curr_top, Inches(11.333), Inches(0.8))
            ttf = tbox.text_frame
            ttf.word_wrap = True
            tp = ttf.paragraphs[0]
            tp.text = s["title"] or s["subtitle"] or "Overview"
            tp.font.size = Pt(30)
            tp.font.bold = True
            tp.font.color.rgb = palette["title"]
            curr_top += Inches(0.85)

            # Subtitle / Summary text
            if s["subtitle"] and s["title"] and s["subtitle"] != s["title"]:
                sbox = slide.shapes.add_textbox(Inches(1.0), curr_top, Inches(11.333), Inches(0.5))
                stf = sbox.text_frame
                stf.word_wrap = True
                sp = stf.paragraphs[0]
                sp.text = s["subtitle"]
                sp.font.size = Pt(15)
                sp.font.color.rgb = palette["muted"]
                curr_top += Inches(0.6)

            # Render Metrics Grid
            if s["metrics"]:
                count = len(s["metrics"])
                cols = min(count, 4)
                card_w = Inches((11.333 - (cols - 1) * 0.3) / cols)
                card_h = Inches(2.8)

                for m_idx, m in enumerate(s["metrics"][:4]):
                    col = m_idx % cols
                    left = Inches(1.0) + col * (card_w + Inches(0.3))

                    c_shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, curr_top, card_w, card_h)
                    c_shape.fill.solid()
                    c_shape.fill.fore_color.rgb = palette["card_bg"]
                    c_shape.line.color.rgb = palette["card_border"]

                    tf = c_shape.text_frame
                    tf.word_wrap = True

                    vp = tf.paragraphs[0]
                    vp.text = m["value"]
                    vp.font.size = Pt(36)
                    vp.font.bold = True
                    vp.font.color.rgb = palette["accent"]

                    lp = tf.add_paragraph()
                    lp.text = m["label"]
                    lp.font.size = Pt(14)
                    lp.font.bold = True
                    lp.font.color.rgb = palette["title"]

                    if m["desc"]:
                        dp = tf.add_paragraph()
                        dp.text = m["desc"]
                        dp.font.size = Pt(11)
                        dp.font.color.rgb = palette["muted"]

            # Render Cards Grid
            elif s["cards"]:
                count = len(s["cards"])
                cols = min(count, 3)
                rows = (count + cols - 1) // cols
                card_w = Inches((11.333 - (cols - 1) * 0.3) / cols)
                card_h = Inches(3.2 / rows)

                for c_idx, c in enumerate(s["cards"]):
                    col = c_idx % cols
                    row = c_idx // cols
                    left = Inches(1.0) + col * (card_w + Inches(0.3))
                    top = curr_top + row * (card_h + Inches(0.25))

                    c_shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, card_w, card_h)
                    c_shape.fill.solid()
                    c_shape.fill.fore_color.rgb = palette["card_bg"]
                    c_shape.line.color.rgb = palette["card_border"]

                    tf = c_shape.text_frame
                    tf.word_wrap = True

                    tp = tf.paragraphs[0]
                    tp.text = c["title"]
                    tp.font.size = Pt(15)
                    tp.font.bold = True
                    tp.font.color.rgb = palette["title"]

                    if c["desc"]:
                        dp = tf.add_paragraph()
                        dp.text = c["desc"]
                        dp.font.size = Pt(12)
                        dp.font.color.rgb = palette["body"]

            # Render Flow Steps
            elif s["steps"]:
                count = min(len(s["steps"]), 4)
                step_w = Inches((11.333 - (count - 1) * 0.25) / count)
                step_h = Inches(3.2)

                for st_idx, st in enumerate(s["steps"][:count]):
                    left = Inches(1.0) + st_idx * (step_w + Inches(0.25))

                    c_shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, curr_top, step_w, step_h)
                    c_shape.fill.solid()
                    c_shape.fill.fore_color.rgb = palette["card_bg"]
                    c_shape.line.color.rgb = palette["card_border"]

                    tf = c_shape.text_frame
                    tf.word_wrap = True

                    np = tf.paragraphs[0]
                    np.text = f"STEP {st['num'] or (st_idx + 1)}"
                    np.font.size = Pt(12)
                    np.font.bold = True
                    np.font.color.rgb = palette["accent"]

                    tp = tf.add_paragraph()
                    tp.text = st["title"]
                    tp.font.size = Pt(15)
                    tp.font.bold = True
                    tp.font.color.rgb = palette["title"]

                    if st["desc"]:
                        dp = tf.add_paragraph()
                        dp.text = st["desc"]
                        dp.font.size = Pt(11)
                        dp.font.color.rgb = palette["body"]

            # Render Bullet Points or Paragraphs
            elif s["bullets"]:
                bbox = slide.shapes.add_textbox(Inches(1.0), curr_top, Inches(11.333), Inches(4.5))
                btf = bbox.text_frame
                btf.word_wrap = True

                for b_idx, bullet in enumerate(s["bullets"]):
                    p = btf.paragraphs[0] if b_idx == 0 else btf.add_paragraph()
                    p.text = f"•   {bullet}"
                    p.font.size = Pt(16)
                    p.font.color.rgb = palette["body"]

            elif s["paragraphs"]:
                pbox = slide.shapes.add_textbox(Inches(1.0), curr_top, Inches(11.333), Inches(4.5))
                ptf = pbox.text_frame
                ptf.word_wrap = True

                for p_idx, para in enumerate(s["paragraphs"]):
                    p = ptf.paragraphs[0] if p_idx == 0 else ptf.add_paragraph()
                    p.text = para
                    p.font.size = Pt(16)
                    p.font.color.rgb = palette["body"]

    os.makedirs(os.path.dirname(os.path.abspath(out_pptx_path)), exist_ok=True)
    prs.save(out_pptx_path)
    return out_pptx_path


def main():
    if len(sys.argv) < 3:
        print("Usage: python3 export-pptx.py <deck-dir-or-index.html> <out.pptx>", file=sys.stderr)
        sys.exit(1)

    target_input = sys.argv[1]
    out_file = sys.argv[2]

    index_html = target_input
    script_dir = os.path.dirname(os.path.abspath(__file__))
    compile_script = os.path.join(script_dir, "compile-markdown.mjs")

    if os.path.isdir(target_input):
        index_html = os.path.join(target_input, "index.html")
        if not os.path.exists(index_html):
            for candidate in ["deck.md", "goal.json"]:
                c_path = os.path.join(target_input, candidate)
                if os.path.exists(c_path):
                    import subprocess
                    subprocess.run(["node", "-e", f"""
                        import('{compile_script}').then(m => {{
                            import('node:fs').then(fs => {{
                                const html = m.compileDeck('{c_path}');
                                fs.writeFileSync('{index_html}', html, 'utf8');
                            }});
                        }});
                    """], check=True)
                    break
    elif target_input.endswith(".md") or target_input.endswith(".json"):
        index_html = os.path.join(os.path.dirname(os.path.abspath(target_input)), "index.html")
        import subprocess
        subprocess.run(["node", "-e", f"""
            import('{compile_script}').then(m => {{
                import('node:fs').then(fs => {{
                    const html = m.compileDeck('{target_input}');
                    fs.writeFileSync('{index_html}', html, 'utf8');
                }});
            }});
        """], check=True)

    if not os.path.exists(index_html):
        print(f"Error: index.html not found at: {index_html}", file=sys.stderr)
        sys.exit(1)

    try:
        exported = export_deck_to_pptx(index_html, out_file)
        print(f"[✓] Successfully exported editable PPTX: {exported}")
    except Exception as e:
        print(f"Export error: {e}", file=sys.stderr)
        sys.exit(1)


if __name__ == "__main__":
    main()
